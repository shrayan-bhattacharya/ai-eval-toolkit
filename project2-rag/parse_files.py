"""
Project 2 - RAG QA Chatbot
File parsing utilities: PDF, Excel, DOCX, TXT
"""

import io
import pdfplumber
import openpyxl
import docx


def parse_pdf(file_bytes: bytes, filename: str) -> str:
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        pages = [page.extract_text() or "" for page in pdf.pages]
    return "\n".join(pages)


def parse_excel(file_bytes: bytes, filename: str) -> str:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    all_rows = []
    for sheet in wb.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(h) if h is not None else "" for h in rows[0]]
        for row in rows[1:]:
            parts = [
                f"{headers[i]}: {row[i]}"
                for i in range(len(row))
                if row[i] is not None and headers[i]
            ]
            if parts:
                all_rows.append(", ".join(parts))
    return "\n".join(all_rows)


def parse_docx(file_bytes: bytes, filename: str) -> str:
    doc = docx.Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)


def parse_file(uploaded_file) -> str:
    """
    Takes a Streamlit UploadedFile, detects type from extension, and returns extracted text.
    Raises ValueError for unsupported formats.
    """
    name = uploaded_file.name.lower()
    file_bytes = uploaded_file.read()

    if name.endswith(".pdf"):
        return parse_pdf(file_bytes, name)
    elif name.endswith((".xlsx", ".xls")):
        return parse_excel(file_bytes, name)
    elif name.endswith(".docx"):
        return parse_docx(file_bytes, name)
    elif name.endswith(".txt"):
        return file_bytes.decode("utf-8")
    else:
        raise ValueError(f"Unsupported file format: {uploaded_file.name}")
